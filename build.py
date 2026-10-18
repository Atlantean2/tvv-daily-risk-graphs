"""Usage: python3 build.py YYYY-MM-DD <graphs dir | risk graphs .pptx>
Renders issues/DATE.html (graphs only) from the dir's 'Graphs DATE.json' (file/title/caption/source),
or from a pptx (largest picture per slide + https links in notes). Always regenerates index.html tabs."""
import sys,os,re,json,glob,datetime
CWD0=os.getcwd()
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)));os.chdir(os.path.dirname(os.path.abspath(__file__)))
from common import *
T="TechVestVentures Daily Risk Graphs"
SKIP=re.compile(r"AIS|PortWatch|transit",re.I)
def render(D,src):
    src=os.path.abspath(src) if os.path.isabs(src) else os.path.join(CWD0,src)
    od=f"issues/img/{D}";os.makedirs(od,exist_ok=True)
    for f in glob.glob(od+"/*"):os.remove(f)
    items=[];skipped=[]
    if os.path.isdir(src):
        js=[j for j in glob.glob(f"{src}/Graphs {D}*.json") if "one-per" not in j] or glob.glob(f"{src}/Graphs*.json")
        for g in json.load(open(js[0]))["graphs"]:
            p=os.path.join(src,g["file"])
            if "selftest" in p or SKIP.search(g["file"]+g.get("title","")): skipped.append(g["file"]);continue
            if not os.path.exists(p) or not str(g.get("source","")).startswith("https://"): skipped.append(g["file"]+" (missing file/source)");continue
            items.append((p,g.get("title",""),g.get("caption",""),[g["source"]]))
    else:
        from pptx import Presentation
        for s in Presentation(src).slides:
            pics=[sh for sh in s.shapes if sh.shape_type==13]
            if not pics:continue
            pic=max(pics,key=lambda x:x.width*x.height)
            n=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
            urls=re.findall(r"https://[^\s)\]>,]+",n)
            txt=[sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            title=s.shapes.title.text if s.shapes.title else (txt[0] if txt else "")
            if SKIP.search(title+n) or not urls: skipped.append(title or "untitled slide");continue
            cap=next((t for t in txt if t!=title and "http" not in t),"")
            items.append((pic.image.blob,title,cap.split("\n")[0],urls[:1]))
    body=[f'<h2 class="desk">Risk graphs — {datetime.date.fromisoformat(D).strftime("%A, %B %-d, %Y")}</h2>']
    for k,(p,t,c,u) in enumerate(items,1):
        fn=f"g{k:02d}.webp";save_img(p,f"{od}/{fn}");body.append(figure(f"img/{D}/{fn}",t,c,u))
    open(f"issues/{D}.html","w").write(page(T,D,"\n".join(body)));print(D,len(items),"graphs; skipped:",skipped)
if __name__=="__main__":
    if len(sys.argv)==3: render(sys.argv[1],sys.argv[2])
    ds=sorted([f[:-5] for f in os.listdir("issues") if re.fullmatch(r"\d{4}-\d{2}-\d{2}\.html",f)],reverse=True)
    index(T,[(d,datetime.date.fromisoformat(d).strftime("%a %b %-d, %Y")) for d in ds],"Graph dates",cur="graphs"); print("tabs:",ds)
